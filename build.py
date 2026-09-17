#!/usr/bin/env python3
"""Bake the public Selby stock-clearance page.

Three gitignored sources sit next to this script (bin locations, costs and
prices never reach the page or the public repo):

  selby-inventory.xlsx   The emailed Selby clearance list (Bin Location,
                         Item Code, Item Description, System Qty). Defines
                         WHICH Selby items are offered — nothing is ever
                         added from the on-hand file (Micheal, 17 Sep 2026).
  selby-on-hand.xlsx     Live-Xero stock-on-hand export ('On hand' sheet:
                         Bin, Code, Description, On hand, ...). Offered
                         items missing from it are REMOVED; kept items take
                         its current quantity.
  x-stock.xlsx           The second stock list (rows marked with an X on
                         the page). ONLY Sheet1 (the curated list with a
                         typed QTY column) is used. Sheet3, a Sage item
                         dump, is deliberately ignored: its only quantity
                         signal is Stock Value / Cost and that proved wrong
                         in reality (640084 derived as 370 on hand —
                         "impossible", Micheal, 17 Sep 2026).

Refreshing: copy the new export(s) over the file(s) above, then

    python3 build.py
    git add index.html build.py README.md template.html   # explicit paths
    git commit -m "stock refresh" -- index.html ...
    # do NOT push — hand the commit to the deploy session

`master-stock.csv` from the scheduler only enriches rows with SX codes,
departments (glass-type filter) and discontinued tags.
"""
import base64
import csv
import datetime
import json
import os

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SELBY_LIST = os.path.join(HERE, "selby-inventory.xlsx")
SELBY_ON_HAND = os.path.join(HERE, "selby-on-hand.xlsx")
X_LIST = os.path.join(HERE, "x-stock.xlsx")
MASTER = os.path.join(HERE, "..", "fitment-scheduler", "catalogue", "master-stock.csv")
LOGO = os.path.join(HERE, "..", "fitment-scheduler", "public", "myglass-logo.png")

# Not offered to other centres (Micheal, 2026-07-22) — survives list refreshes
EXCLUDE_CODES = {"104270", "19-0422", "WT01"}
EXCLUDE_DESC_PREFIXES = ("WB FLAT", "WB STD")
# Tools/consumables stay off the page (same call as the July removals; the
# marked list's Sage dump mixes them in with the glass). Vehicle descriptions never
# contain these — "WIPER HOLE" / "TOOL" never collide because of the
# surrounding words checked here.
EXCLUDE_DESC_WORDS = (" TOOL", "TOOL(", "SCRAPER", "ADHESIVE", "PRIMER",
                      "DOUBLE SIDED TAPE", "RESIN", "KNIFE", "PULLER",
                      "WIRE FEEDER", "WYNNS", "DRILL", "WIPER BLADE",
                      "MASKING TAPE", "BASEBALL CAP", "BUCKET HAT",
                      "SUCTION CUP", "PIANO WIRE", "SPATULA",
                      "POLYURETHANE GUN", "SENSOR TACK", "RAINSENSOR SHEET",
                      "FITMENT KIT")
# Catalogue cross-reference placeholders ("REFER TO 070529EM"), not stock
EXCLUDE_DESC_START = ("REFER TO ",)


def norm_code(v):
    """Sage/Excel hand numeric codes back as ints or floats; 1028.0 -> '1028'."""
    if v is None:
        return ""
    return str(int(v)) if isinstance(v, (int, float)) else str(v).strip()


def excluded(code, desc):
    d = desc.upper()
    return (code in EXCLUDE_CODES or d.startswith(EXCLUDE_DESC_PREFIXES)
            or d.startswith(EXCLUDE_DESC_START)
            or any(w in d for w in EXCLUDE_DESC_WORDS))


def selby_on_hand():
    """{code(lower) -> qty} from the live-Xero export."""
    ws = openpyxl.load_workbook(SELBY_ON_HAND, data_only=True)["On hand"]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = norm_code(row[1])
        if not code:
            continue
        try:
            out[code.lower()] = out.get(code.lower(), 0) + int(float(row[3] or 0))
        except (TypeError, ValueError):
            pass
    return out


def selby_items():
    """Offered list filtered to what is actually on hand today."""
    on_hand = selby_on_hand()
    ws = openpyxl.load_workbook(SELBY_LIST, data_only=True).worksheets[0]
    header = [str(c or "").strip().lower()
              for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col = {n: header.index(n) for n in ("item code", "item description", "system qty")}
    kept, removed, seen = [], [], set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        code = norm_code(row[col["item code"]])
        desc = str(row[col["item description"]] or "").strip()
        try:
            listed_qty = int(float(row[col["system qty"]] or 0))
        except (TypeError, ValueError):
            listed_qty = 0
        if not code or not desc or listed_qty <= 0 or excluded(code, desc):
            continue
        # The same code can sit in two bins = two list rows; the on-hand
        # figure is already the code's total, so one row carries it
        if code.lower() in seen:
            continue
        seen.add(code.lower())
        qty_now = on_hand.get(code.lower(), 0)
        if qty_now <= 0:
            removed.append({"code": code, "desc": desc, "qty": listed_qty})
        else:
            kept.append({"code": code, "desc": desc, "qty": qty_now,
                         "listed_qty": listed_qty})
    return kept, removed


def x_items(selby_on_hand_codes):
    """Sheet1 of the marked list (typed QTY column only — see the module
    docstring for why Sheet3 is ignored). Only codes Selby does NOT hold
    are offered (Micheal, 17 Sep 2026) — no duplicate rows per part."""
    wb = openpyxl.load_workbook(X_LIST, data_only=True)
    by_code = {}
    for row in wb["Sheet1"].iter_rows(min_row=2, values_only=True):
        code = norm_code(row[0])
        desc = str(row[1] or "").strip()
        try:
            qty = int(float(row[2] or 0))
        except (TypeError, ValueError):
            qty = 0
        if code and desc and qty > 0 and not excluded(code, desc):
            by_code[code.lower()] = {"code": code, "desc": desc, "qty": qty}
    return [i for c, i in by_code.items() if c not in selby_on_hand_codes]


def main():
    master = {r["wd_code"].strip().lower(): r
              for r in csv.DictReader(open(MASTER, newline="", encoding="utf-8-sig"))}

    def enrich(item, marked):
        m = master.get(item["code"].lower(), {})
        return {
            "code": item["code"],
            "sx": (m.get("sx_code") or "").strip(),
            "desc": item["desc"],
            "qty": item["qty"],
            "dept": (m.get("department") or "").strip(),
            "disc": (m.get("discontinued") or "").strip() in ("1", "yes", "true"),
            "x": marked,
        }

    selby, removed = selby_items()
    marked = x_items({i["code"].lower() for i in selby} |
                     {c for c, q in selby_on_hand().items() if q > 0})
    items = [enrich(i, False) for i in selby] + [enrich(i, True) for i in marked]
    items.sort(key=lambda i: (i["desc"], i["x"]))

    logo = base64.b64encode(open(LOGO, "rb").read()).decode()
    newest = max(os.path.getmtime(p) for p in (SELBY_ON_HAND, X_LIST, SELBY_LIST))
    updated = datetime.date.fromtimestamp(newest).strftime("%-d %B %Y")

    tpl = open(os.path.join(HERE, "template.html"), encoding="utf-8").read()
    html = (tpl
            .replace("__ITEMS__", json.dumps(items, separators=(",", ":")))
            .replace("__LOGO__", logo)
            .replace("__UPDATED__", updated)
            .replace("__COUNT__", str(len(items))))
    with open(os.path.join(HERE, "index.html"), "w", encoding="utf-8") as f:
        f.write(html)
    print(f"index.html written: {len(items)} items "
          f"({len(selby)} Selby + {len(marked)} marked), list dated {updated}")
    print(f"Selby items removed (not on hand any more): {len(removed)}")
    # For the removal report
    with open(os.path.join(HERE, "last-removed.json"), "w", encoding="utf-8") as f:
        json.dump(removed, f, indent=1)


if __name__ == "__main__":
    main()
